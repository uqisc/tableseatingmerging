"""
Wedding Seating Allocator
==========================
Reads:
  1. A guest-list CSV (ticketing export) containing First name, Last name,
     Order id, and Dietary requirements per attendee.
  2. A table-preferences XLSX (form export) where each row is one group
     submission: "How many people are in your group?" (Group of 5 / 10)
     plus up to 10 "Attendee N: Name / Order ID / Phone Number" slots.

Logic:
  - For each submission, count how many attendee slots are actually filled
    (ignoring blanks like "Nil", "N/a", "None", empty cells).
  - If that count matches the declared group size (5 or 10) -> the group is
    seat-able. Tables physically seat 10, so:
      - a "Group of 10" gets a whole table to itself.
      - two "Group of 5" submissions are paired together (in the order
        they appear) to share one table of 10. A leftover, unpaired group
        of 5 (odd one out) still gets seated on its own table - just half
        empty - rather than being held back.
  - If the headcount does NOT match the declared size -> every valid name
    in that group is set aside on a "Needs Review" list instead of being
    seated (no table number).
  - Each valid attendee is matched to the guest-list CSV in this priority
    order: Order ID (+ Name, since many order IDs cover multiple people)
    -> full Name across the whole guest list -> Phone Number. Each step is
    only tried if the previous one didn't land on exactly one confident
    row. Matching only ever happens for people who actually appear in the
    preferences submissions - the guest list is never searched on its own.
  - Anyone who still can't be matched with confidence after all three
    steps means their WHOLE submission (the entire group) is held for
    review, not just that one person - a group is only seated if every
    member of it is confidently matched AND the headcount matches what
    was declared.

Output:
  A single .xlsx with three sheets:
    - "Seating Plan": First name | Last name | Table number | Dietary | Submitter email
    - "Needs Review": First name | Last name | Order ID | Dietary | Submitter email | Reason
      Rows are highlighted where that specific person is the actual reason
      the group is held back (e.g. their Order ID/name/phone didn't match
      anything) - other rows in the same group are shown plain, since
      they're only there because the group is kept together, not because
      something's wrong with them individually. A group-size mismatch has
      no single culprit, so no row is highlighted in that case.
    - "No Submission": First name | Last name | Order ID | Dietary
      (guests on the guest list who were never referenced - by Order ID,
      name, or phone - anywhere in the preferences file)

USAGE
-----
1. Edit the three settings below (GUEST_LIST_CSV, PREFERENCES_XLSX,
   OUTPUT_XLSX) to point at your files.
2. Run:  python allocate_seating.py
"""

import re
import pandas as pd
import openpyxl

# ----------------------------------------------------------------------
# SETTINGS - edit these paths
# ----------------------------------------------------------------------
GUEST_LIST_CSV = "guest_list.csv"
PREFERENCES_XLSX = "table_preferences.xlsx"
OUTPUT_XLSX = "seating_output.xlsx"

TOTAL_TABLES = 95

# Values that mean "this attendee slot is actually empty"
BLANK_VALUES = {"", "nil", "n/a", "na", "none", "-", "tbc", "tba"}


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------
def is_blank(value) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text.lower() in BLANK_VALUES


def clean_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def norm_order_id(value) -> str:
    return clean_str(value).upper()


def norm_name(value) -> str:
    # lowercase, strip apostrophes/hyphens/punctuation, collapse whitespace
    text = clean_str(value).lower()
    text = re.sub(r"[’'`\-.]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def norm_phone(value) -> str:
    """Keep only digits, then use the last 9 (AU mobile local number) so
    '0450727505', '+61450727505' and '61450727505' all compare equal."""
    digits = re.sub(r"\D", "", clean_str(value))
    return digits[-9:] if len(digits) >= 9 else digits


def parse_group_size(text) -> int:
    """'Group of 5' / 'Group of 10' -> 5 / 10. Returns None if unparseable."""
    if text is None:
        return None
    match = re.search(r"\d+", str(text))
    return int(match.group()) if match else None


def build_dietary(row) -> str:
    """Combine 'Dietary requirements' + the free-text 'other' field."""
    main = clean_str(row.get("Dietary requirements"))
    other = clean_str(row.get("Dietary requirements other"))
    if main.lower() in ("", "nan"):
        main = ""
    if other.lower() in ("", "nan"):
        other = ""

    if not main and not other:
        return ""
    if "other" in main.lower() and other:
        # Replace the bare "Other" token with "Other (detail)"
        parts = [p.strip() for p in main.split(",")]
        parts = [f"Other ({other})" if p.lower() == "other" else p for p in parts]
        return ", ".join(parts)
    if other and not main:
        return other
    return main


# ----------------------------------------------------------------------
# Step 1: Load and index the guest list
# ----------------------------------------------------------------------
def load_guest_list(path):
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    df["_order_id_norm"] = df["Order id"].map(norm_order_id)
    df["_name_norm"] = (df["First name"].fillna("") + " " + df["Last name"].fillna("")).map(norm_name)
    df["_phone_norm"] = df["Buyer mobile"].map(norm_phone) if "Buyer mobile" in df.columns else ""
    df["_dietary"] = df.apply(build_dietary, axis=1)

    by_order_id = {}
    by_name = {}
    by_phone = {}
    for _, row in df.iterrows():
        by_order_id.setdefault(row["_order_id_norm"], []).append(row)
        if row["_name_norm"]:
            by_name.setdefault(row["_name_norm"], []).append(row)
        if row["_phone_norm"]:
            by_phone.setdefault(row["_phone_norm"], []).append(row)
    return df, by_order_id, by_name, by_phone


def match_guest(order_id, name, phone, by_order_id, by_name, by_phone):
    """
    Match priority: Order ID -> full Name -> Phone Number.
    Each step is only tried if the previous one failed to land on exactly
    one confident guest-list row. Returns
    (first_name, last_name, dietary, reason_if_any, matched, method, touched_row_ids).
    reason_if_any is None and matched is True on a clean match. touched_row_ids
    is a list of guest-list dataframe indices this lookup touched (the
    matched row on success, or any candidate rows it considered but
    couldn't confidently pick between on failure) - used to track which
    guest-list rows were ever referenced by the preferences file at all,
    for the "No Submission" report.
    When not confident, first/last/dietary are None so the caller falls
    back to the name as typed in the preferences file (never borrows
    another attendee's details).
    """
    # --- 1. Order ID -----------------------------------------------------
    oid = norm_order_id(order_id)
    candidates = by_order_id.get(oid, [])

    if len(candidates) == 1:
        row = candidates[0]
        return row["First name"], row["Last name"], row["_dietary"], None, True, "order id", [row.name]

    if len(candidates) > 1:
        target = norm_name(name)
        exact = [c for c in candidates if c["_name_norm"] == target]
        if len(exact) == 1:
            row = exact[0]
            return row["First name"], row["Last name"], row["_dietary"], None, True, "order id + name", [row.name]

        partial = [c for c in candidates if target and (target in c["_name_norm"] or c["_name_norm"] in target)]
        if len(partial) == 1:
            row = partial[0]
            return row["First name"], row["Last name"], row["_dietary"], None, True, "order id + partial name", [row.name]

    # --- 2. Full name, across the whole guest list ------------------------
    target = norm_name(name)
    name_candidates = by_name.get(target, [])
    if len(name_candidates) == 1:
        row = name_candidates[0]
        return row["First name"], row["Last name"], row["_dietary"], None, True, "name", [row.name]

    # --- 3. Phone number ---------------------------------------------------
    target_phone = norm_phone(phone)
    phone_candidates = by_phone.get(target_phone, []) if target_phone else []
    if len(phone_candidates) == 1:
        row = phone_candidates[0]
        return row["First name"], row["Last name"], row["_dietary"], None, True, "phone", [row.name]

    # --- Nothing worked confidently ----------------------------------------
    if not candidates and not name_candidates and not phone_candidates:
        reason = "Order ID, name and phone number all failed to match the guest list"
        touched_ids = []
    else:
        reason = "Order ID, name and/or phone number match more than one guest - could not confirm, please verify"
        # even though ambiguous, these specific rows WERE referenced by the
        # preferences file, so they shouldn't be flagged as "no submission"
        touched_ids = [c.name for c in candidates] + [c.name for c in name_candidates] + [c.name for c in phone_candidates]
    return None, None, "", reason, False, None, touched_ids


# ----------------------------------------------------------------------
# Step 2: Walk the preferences submissions
# ----------------------------------------------------------------------
def load_submissions(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    header_index = {h: i for i, h in enumerate(headers)}

    submissions = []
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        group_size_text = row[header_index["How many people are in your group?"]]
        group_size = parse_group_size(group_size_text)

        attendees = []
        for n in range(1, 11):
            name_col = f"Attendee {n}: Name"
            oid_col = f"Attendee {n}: Order ID"
            phone_col = f"Attendee {n}: Phone Number"
            if name_col not in header_index:
                break
            name = row[header_index[name_col]]
            oid = row[header_index[oid_col]]
            phone = row[header_index[phone_col]] if phone_col in header_index else None
            if not is_blank(name):
                attendees.append((clean_str(name), clean_str(oid), clean_str(phone)))

        submissions.append({
            "submitter": row[header_index.get("Your Full Name")] if "Your Full Name" in header_index else "",
            "submitter_email": clean_str(row[header_index["Your Email Address"]]) if "Your Email Address" in header_index else "",
            "group_size_declared": group_size,
            "attendees": attendees,
        })
    return submissions


TABLE_CAPACITY = 10  # seats per physical table; a "Group of 5" fills half


# ----------------------------------------------------------------------
# Step 3: Allocate tables and match dietary info
# ----------------------------------------------------------------------
def allocate(submissions, by_order_id, by_name, by_phone, total_tables=TOTAL_TABLES):
    seated_rows = []
    review_rows = []
    match_notes = []  # console-only log of fallback matches, for transparency
    touched_row_ids = set()  # every guest-list row referenced anywhere in the preferences file

    # ---- Pass 1: match everyone and decide which whole groups are seat-able
    seatable_groups = []  # each: {"size": 5 or 10, "attendees": [...], "submitter_email": ...}

    for sub in submissions:
        declared = sub["group_size_declared"]
        attendees = sub["attendees"]
        actual_count = len(attendees)
        submitter_email = sub.get("submitter_email", "")

        size_mismatch = declared is None or actual_count != declared

        matched_attendees = []
        any_unmatched = False
        for name, oid, phone in attendees:
            first, last, dietary, reason, matched, method, row_ids = match_guest(
                oid, name, phone, by_order_id, by_name, by_phone
            )
            touched_row_ids.update(row_ids)
            if not matched:
                any_unmatched = True
                # no confident guest-list match at all - use the name as
                # typed, never borrow another attendee's row
                parts = name.split(" ", 1)
                first = parts[0]
                last = parts[1] if len(parts) > 1 else ""
                dietary = ""
            matched_attendees.append({
                "first": first, "last": last, "oid": oid, "dietary": dietary,
                "reason": reason, "matched": matched, "method": method,
            })

        group_ok = not size_mismatch and not any_unmatched

        if group_ok:
            seatable_groups.append({
                "size": declared,
                "attendees": matched_attendees,
                "submitter_email": submitter_email,
            })
        else:
            # Whole submission goes to review: either the declared group
            # size doesn't match how many people were actually listed, or
            # at least one attendee in the group couldn't be matched at
            # all - in that case the entire group is held back together
            # rather than seating everyone except the unmatched person.
            reasons_for_group = []
            if size_mismatch:
                reasons_for_group.append(
                    f"Group declared '{declared}' but {actual_count} attendee(s) listed - mismatch"
                )
            culprit_names = [
                f"{a['first']} {a['last']}".strip() for a in matched_attendees if not a["matched"]
            ]
            if culprit_names:
                reasons_for_group.append(
                    f"Held for review because {', '.join(culprit_names)} (same group) could not be "
                    f"matched to the guest list"
                )
            base_reason = "; ".join(reasons_for_group)

            for a in matched_attendees:
                if not a["matched"]:
                    full_reason = f"Could not be matched to the guest list - {a['reason']}"
                    if size_mismatch:
                        full_reason += f"; also, {reasons_for_group[0].lower()}"
                else:
                    full_reason = base_reason
                review_rows.append({
                    "First name": a["first"],
                    "Last name": a["last"],
                    "Order ID": a["oid"],
                    "Dietary": a["dietary"],
                    "Submitter email": submitter_email,
                    "Reason": full_reason,
                    "_flag": not a["matched"],  # this specific person is the actual culprit
                })

    # ---- Pass 2: assign tables. Groups of 10 get a table to themselves;
    # groups of 5 are paired together (in the order they appear) so two of
    # them share one 10-seat table. A leftover unpaired group of 5 still
    # gets seated - just on a half-empty table - rather than being held
    # back, since the people themselves are all valid and confirmed.
    next_table = 1
    tables_used = 0
    pending_five = None
    half_empty_tables = []

    def take_table():
        nonlocal next_table, tables_used
        if next_table > total_tables:
            return None
        t = next_table
        next_table += 1
        tables_used += 1
        return t

    def seat_group(group, table_number):
        for a in group["attendees"]:
            if a["method"] != "order id":
                match_notes.append(f"{a['first']} {a['last']} (order {a['oid']}) matched via {a['method']}")
            seated_rows.append({
                "First name": a["first"],
                "Last name": a["last"],
                "Table number": table_number,
                "Dietary": a["dietary"],
                "Submitter email": group["submitter_email"],
            })

    def group_to_review(group, reason):
        for a in group["attendees"]:
            review_rows.append({
                "First name": a["first"],
                "Last name": a["last"],
                "Order ID": a["oid"],
                "Dietary": a["dietary"],
                "Submitter email": group["submitter_email"],
                "Reason": reason,
                "_flag": False,  # systemic issue (e.g. ran out of tables), not this person's fault
            })

    for group in seatable_groups:
        if group["size"] == TABLE_CAPACITY:
            table_number = take_table()
            if table_number is None:
                group_to_review(group, "Ran out of tables - please assign manually")
            else:
                seat_group(group, table_number)
        else:
            # group of 5 (or any size smaller than a full table)
            if pending_five is None:
                pending_five = group
            else:
                table_number = take_table()
                if table_number is None:
                    group_to_review(pending_five, "Ran out of tables - please assign manually")
                    group_to_review(group, "Ran out of tables - please assign manually")
                else:
                    seat_group(pending_five, table_number)
                    seat_group(group, table_number)
                pending_five = None

    if pending_five is not None:
        # odd one out with no partner group of 5 to pair with
        table_number = take_table()
        if table_number is None:
            group_to_review(pending_five, "Ran out of tables - please assign manually")
        else:
            seat_group(pending_five, table_number)
            half_empty_tables.append(table_number)

    if half_empty_tables:
        match_notes.append(
            f"Table(s) {', '.join(str(t) for t in half_empty_tables)} seated with only one group of 5 "
            f"(no second group of 5 left to pair with) - half-empty, seats 5 of 10."
        )

    return seated_rows, review_rows, tables_used, match_notes, touched_row_ids


def find_no_submission_guests(guest_df, touched_row_ids):
    """Guest-list rows never referenced anywhere in the preferences file."""
    no_submission_rows = []
    for idx, row in guest_df.iterrows():
        if idx not in touched_row_ids:
            no_submission_rows.append({
                "First name": row["First name"],
                "Last name": row["Last name"],
                "Order ID": row["Order id"],
                "Dietary": row["_dietary"],
            })
    return no_submission_rows


# ----------------------------------------------------------------------
# Step 4: Write output workbook
# ----------------------------------------------------------------------
def write_output(seated_rows, review_rows, no_submission_rows, path):
    seated_df = pd.DataFrame(
        seated_rows, columns=["First name", "Last name", "Table number", "Dietary", "Submitter email"]
    )
    seated_df = seated_df.sort_values(["Table number", "Last name"]).reset_index(drop=True)

    review_df = pd.DataFrame(
        review_rows, columns=["First name", "Last name", "Order ID", "Dietary", "Submitter email", "Reason"]
    )

    no_submission_df = pd.DataFrame(
        no_submission_rows, columns=["First name", "Last name", "Order ID", "Dietary"]
    ).sort_values(["Last name", "First name"]).reset_index(drop=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        seated_df.to_excel(writer, sheet_name="Seating Plan", index=False)
        review_df.to_excel(writer, sheet_name="Needs Review", index=False)
        no_submission_df.to_excel(writer, sheet_name="No Submission", index=False)

    # light formatting pass: bold headers, autosize columns
    wb = openpyxl.load_workbook(path)
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        for col_cells in sheet.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            sheet.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    # highlight the specific attendee(s) actually responsible for a review
    # flag (not just everyone in a group held back on their behalf)
    review_sheet = wb["Needs Review"]
    flag_fill = openpyxl.styles.PatternFill(start_color="FFF6A0", end_color="FFF6A0", fill_type="solid")
    flags = [r.get("_flag", False) for r in review_rows]
    for i, flagged in enumerate(flags):
        if flagged:
            excel_row = i + 2  # +1 for header, +1 for 1-indexing
            for cell in review_sheet[excel_row]:
                cell.fill = flag_fill
    # legend, two rows below the data
    legend_row = len(review_rows) + 3
    legend_cell = review_sheet.cell(row=legend_row, column=1, value="Highlighted = this specific person is why the group needs review")
    legend_cell.fill = flag_fill
    legend_cell.font = openpyxl.styles.Font(italic=True, size=9)

    wb.save(path)


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------
def main():
    guest_df, by_order_id, by_name, by_phone = load_guest_list(GUEST_LIST_CSV)
    submissions = load_submissions(PREFERENCES_XLSX)
    seated_rows, review_rows, tables_used, match_notes, touched_row_ids = allocate(
        submissions, by_order_id, by_name, by_phone
    )
    no_submission_rows = find_no_submission_guests(guest_df, touched_row_ids)
    write_output(seated_rows, review_rows, no_submission_rows, OUTPUT_XLSX)

    print(f"Submissions processed : {len(submissions)}")
    print(f"Tables allocated      : {tables_used} / {TOTAL_TABLES}")
    print(f"Guests seated         : {len(seated_rows)}")
    print(f"Guests needing review : {len(review_rows)}")
    print(f"Guests with no submission : {len(no_submission_rows)}")
    print(f"Output written to     : {OUTPUT_XLSX}")
    if match_notes:
        print(f"\n{len(match_notes)} guest(s) matched via name/phone fallback instead of order ID:")
        for note in match_notes:
            print(f"  - {note}")


if __name__ == "__main__":
    main()
